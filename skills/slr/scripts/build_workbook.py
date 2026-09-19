from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule

F = "Arial"
HDR = PatternFill("solid", fgColor="1F3864")
INP = PatternFill("solid", fgColor="FFF2CC")   # input cells
CALC = PatternFill("solid", fgColor="E2EFDA")  # formula cells
EXF = PatternFill("solid", fgColor="EDEDED")   # example row
thin = Side(style="thin", color="BFBFBF")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
N = 1000  # last data row

wb = Workbook()

def header(ws, row, cols, widths=None):
    for i, c in enumerate(cols, 1):
        x = ws.cell(row=row, column=i, value=c)
        x.font = Font(name=F, bold=True, color="FFFFFF", size=10)
        x.fill = HDR
        x.alignment = Alignment(wrap_text=True, vertical="center")
        x.border = BOX
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
    ws.freeze_panes = ws.cell(row=row + 1, column=3)

def style_all(ws):
    for r in ws.iter_rows():
        for c in r:
            if c.font is None or c.font.name != F:
                c.font = Font(name=F, size=10, bold=c.font.bold if c.font else False,
                              color=c.font.color if c.font and c.font.color else None,
                              italic=c.font.italic if c.font else False)

def title(ws, text, sub=None):
    ws["A1"] = text
    ws["A1"].font = Font(name=F, bold=True, size=14)
    if sub:
        ws["A2"] = sub
        ws["A2"].font = Font(name=F, italic=True, size=10, color="595959")

# ---------------------------------------------------------------- README
ws = wb.active
ws.title = "README"
title(ws, "SLR screening workbook (Wolfswinkel 5 stages + PRISMA 2020)",
      "Template for the `slr` skill. Yellow = you type. Green = formula, do not overwrite. Grey row = example, delete before use.")
rows = [
    ("Sheet", "Purpose", "Stage (Wolfswinkel / PRISMA)"),
    ("Protocol", "RQ, sub-RQs, inclusion/exclusion criteria with justification, sources", "1 Define (steps 1.1 to 1.4)"),
    ("Search_Log", "One row per database search: date, exact string, filters, hits, filter-removed records", "2 Search / PRISMA Identification"),
    ("Records", "Every exported record after import; duplicate flag; document-level filter check", "3 Select / PRISMA Identification"),
    ("TA_Screening", "Title, abstract, keyword screening by two reviewers; agreement and final decision", "3 Select / PRISMA Screening"),
    ("FT_Screening", "Full-text eligibility by two reviewers; coded exclusion reasons", "3 Select / PRISMA Eligibility"),
    ("Reliability", "Percent agreement (Wolfswinkel minimum 90%) and Cohen's kappa for both rounds", "3 Select (inter-coder check)"),
    ("PRISMA_Counts", "All flow numbers by formula, plus arithmetic cross-checks", "PRISMA 2020 flow"),
    ("PRISMA_Flow", "Flow diagram drawn in cells and linked to PRISMA_Counts", "PRISMA 2020 flow"),
    ("Evidence_Table", "One row per included study (Authors, Title, Journal, Summary, Details, Why it matters for the RQ)", "4 Analyze (excerpting)"),
    ("Concept_Matrix", "Papers x concepts (Wolfswinkel Tables 3 and 4); merge and split concepts over time", "4 Analyze (open, axial, selective coding)"),
    ("Logbook", "Dated decisions and criteria changes (process diary)", "All stages"),
    ("Lists", "Decision codes and exclusion-reason codes used by drop-downs", "Support"),
]
for i, r in enumerate(rows, 4):
    for j, v in enumerate(r, 1):
        c = ws.cell(row=i, column=j, value=v)
        c.border = BOX
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if i == 4:
            c.font = Font(name=F, bold=True, color="FFFFFF"); c.fill = HDR
ws["A19"] = "Screening codes: Include / Exclude / Maybe. Maybe counts as a disagreement-risk and must be resolved before the final decision."
ws["A20"] = "Final decision rule: if both reviewers agree, that decision stands. If not, discuss; if still unresolved, a third person decides. Record it in the Final column."
ws["A21"] = "Rows 5 onward are counted. The grey example row in row 4 of each screening sheet is NOT counted and can be deleted."
for a in ("A19", "A20", "A21"):
    ws[a].font = Font(name=F, size=10)
ws.column_dimensions["A"].width = 18
ws.column_dimensions["B"].width = 80
ws.column_dimensions["C"].width = 38

# ---------------------------------------------------------------- Lists
ls = wb.create_sheet("Lists")
ls["A1"], ls["B1"], ls["C1"] = "Decision", "Exclusion reason code", "Reason description (edit to fit your criteria)"
for c in ("A1", "B1", "C1"):
    ls[c].font = Font(name=F, bold=True, color="FFFFFF"); ls[c].fill = HDR
for i, v in enumerate(["Include", "Exclude", "Maybe"], 2):
    ls.cell(row=i, column=1, value=v)
reasons = [
    ("R1", "Not empirical (conceptual, opinion, editorial)"),
    ("R2", "Wrong document type (book chapter, conference paper, thesis, review)"),
    ("R3", "Not English"),
    ("R4", "Outside time window"),
    ("R5", "Outside field / subject area"),
    ("R6", "Off topic: does not address the research question"),
    ("R7", "Wrong population or actor (e.g. follower gender instead of leader gender)"),
    ("R8", "Key construct missing (e.g. no gender dimension, no trust outcome)"),
    ("R9", "Below quality threshold (e.g. impact factor)"),
    ("R10", "Full text not retrievable"),
    ("R11", "Other (explain in Notes)"),
]
for i, (a, b) in enumerate(reasons, 2):
    ls.cell(row=i, column=2, value=a); ls.cell(row=i, column=3, value=b)
ls.column_dimensions["A"].width = 12; ls.column_dimensions["B"].width = 22; ls.column_dimensions["C"].width = 70

def dvs():
    return (DataValidation(type="list", formula1="=Lists!$A$2:$A$4", allow_blank=True),
            DataValidation(type="list", formula1="=Lists!$B$2:$B$12", allow_blank=True),
            DataValidation(type="list", formula1='"Y,N"', allow_blank=True))

# ---------------------------------------------------------------- Protocol
p = wb.create_sheet("Protocol", 1)
title(p, "Protocol: define (Wolfswinkel stage 1)")
p["A3"], p["B3"] = "Research question", ""
p["A4"], p["B4"] = "Sub-questions", ""
p["A5"], p["B5"] = "Review type / method", "Grounded-theory literature review (Wolfswinkel et al., 2013); reporting per PRISMA 2020 (Page et al., 2021)"
p["A6"], p["B6"] = "Team and roles", "Reviewer A: ...   Reviewer B: ...   Tie-breaker: ..."
for r in range(3, 7):
    p.cell(row=r, column=1).font = Font(name=F, bold=True)
    p.cell(row=r, column=2).fill = INP
    p.cell(row=r, column=2).alignment = Alignment(wrap_text=True, vertical="top")
header(p, 8, ["Criterion", "Decision", "Justification (one sentence)", "Applied where (DB filter / TA / FT)", "Step (Table 1)"])
crit = [
    ("Document type", "Journal articles only", "Peer-reviewed, legitimised corpus", "DB filter", "1.1"),
    ("Empirical", "Empirical studies only", "Need evidence, not opinion", "TA + FT", "1.1"),
    ("Language", "English", "All reviewers can read and check", "DB filter", "1.1"),
    ("Field", "e.g. Business, Management, Social Sciences", "Where the construct is studied", "DB filter", "1.2"),
    ("Time window", "e.g. 2019 to 2026", "Relevance to current context", "DB filter", "1.1"),
    ("Quality proxy", "e.g. impact factor above 1 (state if NOT applied)", "Screen for outlet quality; report as limitation", "DB filter / TA", "1.1"),
    ("Topic-specific", "Operational definitions of key constructs", "Decide borderline terms in advance", "TA + FT", "1.1"),
]
for i, r in enumerate(crit, 9):
    for j, v in enumerate(r, 1):
        c = p.cell(row=i, column=j, value=v); c.fill = INP; c.border = BOX
        c.alignment = Alignment(wrap_text=True, vertical="top")
p["A17"] = "Sources (step 1.3) and search terms (step 1.4) are recorded in Search_Log."
p["A17"].font = Font(name=F, italic=True)
for col, w in zip("ABCDE", (24, 46, 46, 30, 12)):
    p.column_dimensions[col].width = w
p.freeze_panes = None

# ---------------------------------------------------------------- Search_Log
s = wb.create_sheet("Search_Log")
title(s, "Search log (stage 2). One row per database and run")
header(s, 3, ["Search ID", "Database / interface", "Date run", "Exact search string (verbatim)", "Fields searched",
              "Hits before filters", "Removed by filters (doc type, language, years, subject)", "Hits after filters", "Filters applied", "Run by"],
       [10, 22, 12, 60, 18, 14, 22, 14, 36, 10])
s.freeze_panes = "A4"
ex = ["EX-S1", "Scopus (example)", "2026-01-07", "TITLE-ABS-KEY( gender OR sex ) AND TITLE-ABS-KEY( leader* ) AND TITLE-ABS-KEY( trust )",
      "Title, abstract, keywords", 291, 126, None, "English; article; 2019 to 2026", "AB"]
for j, v in enumerate(ex, 1):
    c = s.cell(row=4, column=j, value=v); c.fill = EXF; c.font = Font(name=F, italic=True, size=10)
for r in range(4, 40):
    h = s.cell(row=r, column=8, value=f'=IF(F{r}="","",F{r}-G{r})'); h.fill = CALC
    for j in (1, 2, 3, 4, 5, 6, 7, 9, 10):
        if r > 4: s.cell(row=r, column=j).fill = INP
s["L3"] = "Example row 4 is not counted (sums start at row 5)."
s["L3"].font = Font(name=F, italic=True, size=9)
s["L4"] = "Other-methods records (citation searching) go in PRISMA_Counts."
s["L4"].font = Font(name=F, italic=True, size=9)

# ---------------------------------------------------------------- Records
rc = wb.create_sheet("Records")
title(rc, "Records: everything exported from the databases (after import)")
header(rc, 3, ["ID", "Source DB", "Authors", "Year", "Title", "Journal", "DOI", "Impact factor / quartile", "Doc type", "Language",
               "Abstract", "Duplicate? (Y/N)", "Passes DB-level criteria? (Y/N)"],
       [8, 12, 26, 7, 50, 26, 22, 12, 12, 10, 60, 12, 16])
rc.freeze_panes = "C4"
ex = ["EX-1", "Scopus", "Belasen, A.T.; Belasen, A.R.", 2024, "Signaling trust during disruptions", "(journal)", "", 3.47, "Article", "English",
      "(abstract text)", "N", "Y"]
for j, v in enumerate(ex, 1):
    c = rc.cell(row=4, column=j, value=v); c.fill = EXF; c.font = Font(name=F, italic=True, size=10)
rc["O3"] = "Keep author strings in one format (Surname, I.; Surname, I.). Add a Year and Journal column: the shared sheet lacked both."
rc["O3"].font = Font(name=F, italic=True, size=9)
_,_,dv_yn=dvs(); rc.add_data_validation(dv_yn); dv_yn.add(f"L5:M{N}")

# ---------------------------------------------------------------- TA_Screening
ta = wb.create_sheet("TA_Screening")
title(ta, "Title / abstract / keyword screening (stage 3, PRISMA 'Screening')",
      "Records that passed DB-level criteria and are not duplicates. Reviewers decide independently BEFORE looking at each other's columns.")
header(ta, 3, ["ID", "Title", "Year", "Reviewer A decision", "Reviewer A reason (code)", "Reviewer B decision", "Reviewer B reason (code)",
               "Agree?", "Final decision", "Final reason (code)", "Resolved by", "Notes / reason for selection"],
       [8, 50, 7, 14, 14, 14, 14, 9, 14, 14, 12, 40])
ta.freeze_panes = "C4"
ex = ["EX-1", "Signaling trust during disruptions", 2024, "Include", "", "Include", "", None, "Include", "", "", "Leader gender and trust both present"]
for j, v in enumerate(ex, 1):
    c = ta.cell(row=4, column=j, value=v); c.fill = EXF; c.font = Font(name=F, italic=True, size=10)
ta["H4"] = '=IF(OR(D4="",F4=""),"",IF(D4=F4,"Yes","No"))'
for r in range(5, N + 1):
    h = ta.cell(row=r, column=8, value=f'=IF(OR(D{r}="",F{r}=""),"",IF(D{r}=F{r},"Yes","No"))'); h.fill = CALC
    for j in (1, 2, 3, 4, 5, 6, 7, 9, 10, 11, 12):
        ta.cell(row=r, column=j).fill = INP
dv_dec,dv_reason,_=dvs(); ta.add_data_validation(dv_dec); ta.add_data_validation(dv_reason)
for col in "DFI": dv_dec.add(f"{col}5:{col}{N}")
for col in "EGJ": dv_reason.add(f"{col}5:{col}{N}")
ta.conditional_formatting.add(f"H5:H{N}", CellIsRule(operator="equal", formula=['"No"'], fill=PatternFill("solid", bgColor="F8CBAD")))

# ---------------------------------------------------------------- FT_Screening
ft = wb.create_sheet("FT_Screening")
title(ft, "Full-text eligibility (stage 3, PRISMA 'Eligibility'). Careful reading",
      "One row per report sought for retrieval: all TA final = Include or Maybe-resolved-to-Include, plus records from citation searching.")
header(ft, 3, ["ID", "Title", "Source (Database / Backward / Forward)", "Retrieved? (Y/N)", "Reviewer A decision", "Reviewer B decision",
               "Agree?", "Final decision", "Exclusion reason (code)", "Resolved by", "Notes / reason for selection"],
       [8, 50, 18, 12, 14, 14, 9, 14, 16, 12, 44])
ft.freeze_panes = "C4"
ex = ["EX-1", "Signaling trust during disruptions", "Database", "Y", "Include", "Include", None, "Include", "", "", "Meets all criteria"]
for j, v in enumerate(ex, 1):
    c = ft.cell(row=4, column=j, value=v); c.fill = EXF; c.font = Font(name=F, italic=True, size=10)
ft["G4"] = '=IF(OR(E4="",F4=""),"",IF(E4=F4,"Yes","No"))'
for r in range(5, N + 1):
    h = ft.cell(row=r, column=7, value=f'=IF(OR(E{r}="",F{r}=""),"",IF(E{r}=F{r},"Yes","No"))'); h.fill = CALC
    for j in (1, 2, 3, 4, 5, 6, 8, 9, 10, 11):
        ft.cell(row=r, column=j).fill = INP
dv_dec,dv_reason,dv_yn=dvs(); ft.add_data_validation(dv_dec); ft.add_data_validation(dv_reason); ft.add_data_validation(dv_yn)
for col in "EFH": dv_dec.add(f"{col}5:{col}{N}")
dv_reason.add(f"I5:I{N}"); dv_yn.add(f"D5:D{N}")
ft.conditional_formatting.add(f"G5:G{N}", CellIsRule(operator="equal", formula=['"No"'], fill=PatternFill("solid", bgColor="F8CBAD")))

# ---------------------------------------------------------------- Reliability
rl = wb.create_sheet("Reliability")
title(rl, "Inter-reviewer reliability", "Wolfswinkel et al. (2013): at least 90% overlap between two reviewers on article selection. Kappa is reported as a complement.")
header(rl, 4, ["Metric", "Title/abstract round", "Full-text round", "Rule / note"], [40, 22, 22, 60])
rl.freeze_panes = None
def ta_rng(c): return f"TA_Screening!${c}$5:${c}${N}"
def ft_rng(c): return f"FT_Screening!${c}$5:${c}${N}"
lab = [
    ("Records with two decisions (N)", f'=COUNTIFS({ta_rng("D")},"<>",{ta_rng("F")},"<>")', f'=COUNTIFS({ft_rng("E")},"<>",{ft_rng("F")},"<>")', "Both reviewers must have decided"),
    ("Agreements", f'=COUNTIF({ta_rng("H")},"Yes")', f'=COUNTIF({ft_rng("G")},"Yes")', ""),
    ("Disagreements", f'=COUNTIF({ta_rng("H")},"No")', f'=COUNTIF({ft_rng("G")},"No")', "Resolve by discussion, then third reviewer"),
    ("Percent agreement", '=IF(B5=0,"",B6/B5)', '=IF(C5=0,"",C6/C5)', "Target: 0.90 or higher"),
    ("Meets 90% standard?", '=IF(B8="","",IF(B8>=0.9,"Yes","NO: recalibrate criteria and re-screen a sample"))',
     '=IF(C8="","",IF(C8>=0.9,"Yes","NO: recalibrate criteria and re-screen a sample"))', "Wolfswinkel et al. (2013), Select stage"),
]
for i, (a, b, c, d) in enumerate(lab, 5):
    rl.cell(row=i, column=1, value=a); rl.cell(row=i, column=2, value=b).fill = CALC
    rl.cell(row=i, column=3, value=c).fill = CALC; rl.cell(row=i, column=4, value=d)
rl["B8"].number_format = "0.0%"; rl["C8"].number_format = "0.0%"
# kappa: categories Include/Exclude/Maybe
rl["A11"] = "Cohen's kappa (3 categories: Include, Exclude, Maybe)"; rl["A11"].font = Font(name=F, bold=True)
def pe(a, b, n):
    parts = [f'(COUNTIF({a},"{k}")*COUNTIF({b},"{k}"))' for k in ("Include", "Exclude", "Maybe")]
    return "(" + "+".join(parts) + f")/({n}^2)"
rl["A12"] = "Expected chance agreement (Pe)"
rl["B12"] = f'=IF(B5=0,"",{pe(ta_rng("D"), ta_rng("F"), "B5")})'
rl["C12"] = f'=IF(C5=0,"",{pe(ft_rng("E"), ft_rng("F"), "C5")})'
rl["A13"] = "Kappa = (Po - Pe) / (1 - Pe)"
rl["B13"] = '=IF(OR(B8="",B12=""),"",IF(B12=1,1,(B8-B12)/(1-B12)))'
rl["C13"] = '=IF(OR(C8="",C12=""),"",IF(C12=1,1,(C8-C12)/(1-C12)))'
rl["D13"] = "Rough guide: 0.61 to 0.80 substantial, above 0.80 almost perfect (Landis & Koch, 1977)"
for c in ("B12", "C12", "B13", "C13"):
    rl[c].fill = CALC; rl[c].number_format = "0.00"

# ---------------------------------------------------------------- PRISMA_Counts
pc = wb.create_sheet("PRISMA_Counts")
title(pc, "PRISMA 2020 counts (all by formula) and cross-checks",
      "Yellow cells are the only manual inputs. Everything else reads the other sheets.")
header(pc, 4, ["Item", "Count", "Source / formula logic", "Check"], [62, 12, 62, 40])
pc.freeze_panes = None
SL = "Search_Log"
rec = "Records"
items = [
    # row 5..
    ("Records identified from databases and registers", f'=SUM({SL}!$F$5:$F$39)', "Sum of hits before filters in Search_Log", ""),                       #5
    ("  of which: number of databases searched", f'=COUNTA({SL}!$B$5:$B$39)', "Rows in Search_Log", ""),                                              #6
    ("Reports from other methods sought for retrieval (backward/forward citation searching)", None, "MANUAL: citation-searched papers that pass your quick relevance look; also list them in FT_Screening", ""),               #7
    ("Records removed before screening: duplicates", f'=COUNTIF({rec}!$L$5:$L${N},"Y")', "Records sheet, Duplicate = Y", ""),                              #8
    ("Records removed before screening: marked ineligible by DB filters", f'=SUM({SL}!$G$5:$G$39)', "Filters applied at search (Search_Log col G)", ""),      #9
    ("Records removed before screening: other reasons", None, "MANUAL (e.g. retracted, automation tool); explain in Logbook", ""),                       #10
    ("Records screened (title/abstract)", f'=COUNTA(TA_Screening!$A$5:$A${N})', "Rows in TA_Screening", '=IF(B11=B5-B8-B9-B10,"OK","MISMATCH: identified minus removed must equal screened")'),  #11
    ("Records excluded at title/abstract", f'=COUNTIF(TA_Screening!$I$5:$I${N},"Exclude")', "TA final decision = Exclude", '=IF(COUNTIF(TA_Screening!$I$5:$I$1000,"Maybe")>0,"Unresolved Maybe remain","OK")'),  #12
    ("Reports sought for retrieval", f'=COUNTA(FT_Screening!$A$5:$A${N})', "Rows in FT_Screening", '=IF(B13=B11-B12+B7,"OK","MISMATCH: screened minus excluded (+ other-method reports) must equal sought")'),  #13
    ("Reports not retrieved", f'=COUNTIF(FT_Screening!$D$5:$D${N},"N")', "Retrieved = N", ""),                                                          #14
    ("Reports assessed for eligibility", '=B13-B14', "Sought minus not retrieved", ""),                                                                  #15
    ("Reports excluded at full text (total)", f'=COUNTIF(FT_Screening!$H$5:$H${N},"Exclude")', "FT final decision = Exclude", '=IF(B16=SUM(B20:B29),"OK","MISMATCH: reason counts do not sum to total excluded")'),  #16
    ("Studies included in review", f'=COUNTIF(FT_Screening!$H$5:$H${N},"Include")', "FT final decision = Include", '=IF(B17=B15-B16,"OK","MISMATCH: assessed minus excluded must equal included")'),  #17
]
for i, (a, b, c, d) in enumerate(items, 5):
    pc.cell(row=i, column=1, value=a)
    x = pc.cell(row=i, column=2, value=b)
    x.fill = INP if b is None else CALC
    pc.cell(row=i, column=3, value=c)
    pc.cell(row=i, column=4, value=d).fill = CALC if d else PatternFill()
pc["B7"] = 0; pc["B10"] = 0
pc["A18"] = "Full-text exclusion reasons (PRISMA requires the reason and count for each)"; pc["A18"].font = Font(name=F, bold=True)
pc["A19"] = "Reason code"; pc["B19"] = "Count"; pc["C19"] = "Description"
for c in ("A19", "B19", "C19"):
    pc[c].font = Font(name=F, bold=True)
codes = [f"R{k}" for k in range(1, 12)]
# rows 20..29 hold R1..R10 in B20:B29 (R11 row 30 listed separately but included in check via 20:30)
for i, code in enumerate(codes, 20):
    pc.cell(row=i, column=1, value=code)
    pc.cell(row=i, column=2, value=f'=COUNTIFS(FT_Screening!$H$5:$H${N},"Exclude",FT_Screening!$I$5:$I${N},A{i})').fill = CALC
    pc.cell(row=i, column=3, value=f'=INDEX(Lists!$C$2:$C$12,MATCH(A{i},Lists!$B$2:$B$12,0))')
pc["D16"] = '=IF(B16=SUM(B20:B30),"OK","MISMATCH: reason counts do not sum to total excluded")'
pc["A32"] = "Global checks"; pc["A32"].font = Font(name=F, bold=True)
pc["A33"] = "Every excluded full-text record has a reason code"
pc["D33"] = f'=IF(COUNTIFS(FT_Screening!$H$5:$H${N},"Exclude",FT_Screening!$I$5:$I${N},"")=0,"OK","Missing reason codes")'
pc["A34"] = "Title/abstract disagreements all resolved (Final decision filled)"
pc["D34"] = f'=IF(COUNTIFS(TA_Screening!$H$5:$H${N},"No",TA_Screening!$I$5:$I${N},"")=0,"OK","Unresolved disagreements")'
pc["A35"] = "Full-text disagreements all resolved"
pc["D35"] = f'=IF(COUNTIFS(FT_Screening!$G$5:$G${N},"No",FT_Screening!$H$5:$H${N},"")=0,"OK","Unresolved disagreements")'
pc["A36"] = "Agreement at least 90% in both rounds"
pc["D36"] = '=IF(AND(Reliability!B9="Yes",Reliability!C9="Yes"),"OK","Check Reliability sheet")'
for r in (33, 34, 35, 36):
    pc.cell(row=r, column=4).fill = CALC
pc.conditional_formatting.add("D5:D36", FormulaRule(formula=['LEFT(D5,8)="MISMATCH"'], fill=PatternFill("solid", bgColor="F8CBAD")))
pc.conditional_formatting.add("D5:D36", FormulaRule(formula=['D5="OK"'], fill=PatternFill("solid", bgColor="C6EFCE")))

# ---------------------------------------------------------------- PRISMA_Flow
fl = wb.create_sheet("PRISMA_Flow")
title(fl, "PRISMA 2020 flow diagram (linked to PRISMA_Counts)")
fl.column_dimensions["A"].width = 16
for col in "BCDEFG":
    fl.column_dimensions[col].width = 20
def box(rng, text, fill="FFFFFF", bold=False):
    fl.merge_cells(rng)
    c = fl[rng.split(":")[0]]
    c.value = text; c.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
    c.font = Font(name=F, size=10, bold=bold)
    first, last = rng.split(":")
    for row in fl[rng]:
        for cell in row:
            cell.border = BOX
            cell.fill = PatternFill("solid", fgColor=fill)
def phase(rng, text):
    fl.merge_cells(rng)
    c = fl[rng.split(":")[0]]
    c.value = text; c.alignment = Alignment(text_rotation=90, horizontal="center", vertical="center")
    c.font = Font(name=F, bold=True, color="FFFFFF");
    for row in fl[rng]:
        for cell in row: cell.fill = HDR; cell.border = BOX
phase("A4:A9", "Identification"); phase("A10:A15", "Screening"); phase("A16:A19", "Included")
box("B4:D6", '="Records identified from databases and registers (n = "&PRISMA_Counts!B5&")"&CHAR(10)&"Databases (n = "&PRISMA_Counts!B6&")"', "DDEBF7")
box("E4:G6", '="Identification via other methods: reports sought from citation searching (n = "&PRISMA_Counts!B7&")"', "DDEBF7")
box("B7:D9", '="Records removed before screening:"&CHAR(10)&"Duplicates (n = "&PRISMA_Counts!B8&")"&CHAR(10)&"Marked ineligible by filters (n = "&PRISMA_Counts!B9&")"&CHAR(10)&"Other reasons (n = "&PRISMA_Counts!B10&")"', "F2F2F2")
box("B10:D11", '="Records screened (n = "&PRISMA_Counts!B11&")"', "DDEBF7")
box("E10:G11", '="Records excluded (n = "&PRISMA_Counts!B12&")"', "F2F2F2")
box("B12:D13", '="Reports sought for retrieval (n = "&PRISMA_Counts!B13&")"', "DDEBF7")
box("E12:G13", '="Reports not retrieved (n = "&PRISMA_Counts!B14&")"', "F2F2F2")
box("B14:D15", '="Reports assessed for eligibility (n = "&PRISMA_Counts!B15&")"', "DDEBF7")
reason_txt = '="Reports excluded (n = "&PRISMA_Counts!B16&")"&' + "&".join(
    [f'IF(PRISMA_Counts!B{r}>0,CHAR(10)&PRISMA_Counts!A{r}&": "&INDEX(Lists!$C$2:$C$12,MATCH(PRISMA_Counts!A{r},Lists!$B$2:$B$12,0))&" (n = "&PRISMA_Counts!B{r}&")","")' for r in range(20, 31)])
box("E14:G15", reason_txt, "F2F2F2")
box("B16:D19", '="Studies included in review (n = "&PRISMA_Counts!B17&")"', "C6EFCE", True)
for r in range(4, 20):
    fl.row_dimensions[r].height = 26
fl.row_dimensions[14].height = 60; fl.row_dimensions[15].height = 60
fl["A21"] = "Copy this block into Word/PowerPoint as a picture, or redraw it as a figure from the counts (Page et al., 2021 template)."
fl["A21"].font = Font(name=F, italic=True, size=9)

# ---------------------------------------------------------------- Evidence_Table
ev = wb.create_sheet("Evidence_Table")
title(ev, "Evidence table: one row per included study (stage 4)")
header(ev, 3, ["ID", "Authors", "Year", "Title", "Journal", "Summary (2 to 4 sentences)", "Context / country", "Design (qual / quant / mixed)",
               "Key results", "Agrees with / differs from other studies", "Limitations / biases", "Why it matters for the RQ"],
       [8, 26, 7, 46, 24, 50, 18, 16, 44, 32, 30, 40])
ev.freeze_panes = "C4"
ex = ["EX-1", "Belasen, A.T.; Belasen, A.R.", 2024, "Signaling trust during disruptions", "(journal)", "(summary)", "USA", "Quantitative",
      "(results)", "Consistent with role congruity findings", "(limitations)", "Shows different trust signals for male and female leaders"]
for j, v in enumerate(ex, 1):
    c = ev.cell(row=4, column=j, value=v); c.fill = EXF; c.font = Font(name=F, italic=True, size=10)
for r in range(5, 60):
    for j in range(1, 13):
        ev.cell(row=r, column=j).fill = INP
        ev.cell(row=r, column=j).alignment = Alignment(wrap_text=True, vertical="top")

# ---------------------------------------------------------------- Concept_Matrix
cm = wb.create_sheet("Concept_Matrix")
title(cm, "Concept matrix (Wolfswinkel Tables 3 and 4): mark X where a paper contributes an excerpt for a concept",
      "Merge, split and relabel concepts as coding proceeds and record each change in the Logbook. Columns after the papers give counts.")
cm["A4"] = "Paper ID"; cm["B4"] = "Authors (year)"
cm["A4"].font = cm["B4"].font = Font(name=F, bold=True, color="FFFFFF"); cm["A4"].fill = cm["B4"].fill = HDR
cm["C3"] = "Category (axial)  ->"; cm["C3"].font = Font(name=F, italic=True)
labels = [("Concept 1 (open code)", "Category A"), ("Concept 2 (open code)", "Category A"), ("Concept 3 (open code)", "Category B"),
          ("Concept 4 (open code)", "Category B"), ("Concept 5 (open code)", "Category C"), ("Concept 6 (open code)", "Category C")]
for j, (cn, cat) in enumerate(labels, 3):
    cm.cell(row=3, column=j, value=cat).font = Font(name=F, italic=True)
    x = cm.cell(row=4, column=j, value=cn); x.font = Font(name=F, bold=True, color="FFFFFF"); x.fill = HDR
    x.alignment = Alignment(wrap_text=True)
    cm.column_dimensions[x.column_letter].width = 16
cm.column_dimensions["A"].width = 10; cm.column_dimensions["B"].width = 30
cm["I4"] = "Concepts per paper"; cm["I4"].font = Font(name=F, bold=True, color="FFFFFF"); cm["I4"].fill = HDR
cm.column_dimensions["I"].width = 18
for r in range(5, 45):
    for j in range(1, 9): cm.cell(row=r, column=j).fill = INP
    cm.cell(row=r, column=9, value=f'=IF(A{r}="","",COUNTIF(C{r}:H{r},"X"))').fill = CALC
cm["A46"] = "Papers per concept"; cm["A46"].font = Font(name=F, bold=True)
for j in range(3, 9):
    col = cm.cell(row=4, column=j).column_letter
    cm.cell(row=46, column=j, value=f'=COUNTIF({col}5:{col}44,"X")').fill = CALC
cm["A5"], cm["B5"], cm["C5"], cm["E5"] = "EX-1", "Belasen & Belasen (2024) [example]", "X", "X"
for c in ("A5", "B5", "C5", "E5"):
    cm[c].font = Font(name=F, italic=True)

# ---------------------------------------------------------------- Logbook
lg = wb.create_sheet("Logbook")
title(lg, "Logbook: process diary (Wolfswinkel: write down the reason for every decision)")
header(lg, 3, ["Date", "Stage / step (Table 1)", "Decision or change", "Reason", "Consequence (what must be redone)", "By"], [12, 20, 50, 44, 44, 8])
lg.freeze_panes = "A4"
ex = ["2026-01-07", "2 Search (2.1)", "Narrowed RQ from 'gender differences' to 'leader gender and follower trust'", "First run returned too many off-topic hits", "Re-run all databases; re-screen", "AB"]
for j, v in enumerate(ex, 1):
    c = lg.cell(row=4, column=j, value=v); c.fill = EXF; c.font = Font(name=F, italic=True, size=10)
for r in range(5, 60):
    for j in range(1, 7): lg.cell(row=r, column=j).fill = INP

for w in wb.worksheets:
    for row in w.iter_rows():
        for c in row:
            f = c.font
            if f is None or f.name != F:
                c.font = Font(name=F, size=f.size if f and f.size and f.size != 11 else 10, bold=f.bold if f else False,
                              italic=f.italic if f else False, color=f.color if f else None)
# sheet order
order = ["README", "Protocol", "Search_Log", "Records", "TA_Screening", "FT_Screening", "Reliability", "PRISMA_Counts",
         "PRISMA_Flow", "Evidence_Table", "Concept_Matrix", "Logbook", "Lists"]
wb._sheets = [wb[n] for n in order]
import sys
wb.save(sys.argv[1] if len(sys.argv) > 1 else "SLR_Screening_Workbook_Template.xlsx")
print("saved")
